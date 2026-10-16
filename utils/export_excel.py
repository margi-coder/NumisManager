from io import BytesIO
from typing import List, Dict, Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

GOLD_FILL = PatternFill(start_color="D4AF37", end_color="D4AF37", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
REGULAR_FONT = Font(name="Calibri", size=10)
BOLD_FONT = Font(name="Calibri", size=10, bold=True)

THIN_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC")
)

def _apply_table_styling(ws, header_row=1):
    """Menerapkan border, alignment, dan penyesuaian lebar kolom otomatis."""
    # Styling Header
    for cell in ws[header_row]:
        cell.fill = GOLD_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    # Styling Content Rows
    for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
        for cell in row:
            cell.font = REGULAR_FONT
            cell.border = THIN_BORDER
            if isinstance(cell.value, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto Width Adjustment
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

def export_stok(items: List[Dict[str, Any]]) -> BytesIO:
    """Ekspor laporan stok koleksi ke file Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Laporan Stok"

    headers = [
        "ID", "Nama Item", "Tahun", "Grade", "Jenis", "Negara",
        "Qty Beli", "Qty Terjual", "Qty Tersisa",
        "HPP Total", "HPP per Pcs", "Harga Jual", "Nilai Stok"
    ]
    ws.append(headers)

    for item in items:
        ws.append([
            item.get("id"),
            item.get("nama_item"),
            item.get("tahun"),
            item.get("grade"),
            item.get("jenis"),
            item.get("negara"),
            item.get("qty", 0),
            item.get("qty_terjual", 0),
            item.get("qty_tersisa", 0),
            item.get("hpp", 0.0),
            item.get("hpp_per_pcs", 0.0),
            item.get("harga_jual", 0.0),
            item.get("nilai_stok", 0.0),
        ])

    _apply_table_styling(ws, header_row=1)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream

def export_penjualan(items: List[Dict[str, Any]]) -> BytesIO:
    """Ekspor laporan penjualan transaksi ke file Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Laporan Penjualan"

    headers = [
        "Kode Jual", "Tanggal Jual", "Nama Item", "Qty Jual",
        "Harga Jual / Pcs", "HPP / Pcs", "Total Jual",
        "Profit / Pcs", "Total Profit", "Margin (%)",
        "Pembeli", "Platform"
    ]
    ws.append(headers)

    for item in items:
        ws.append([
            item.get("kode_jual"),
            item.get("tanggal_jual"),
            item.get("nama_item"),
            item.get("qty_jual", 0),
            item.get("harga_jual_pcs", 0.0),
            item.get("hpp_pcs", 0.0),
            item.get("total_jual", 0.0),
            item.get("profit_pcs", 0.0),
            item.get("total_profit", 0.0),
            item.get("margin_pct", 0.0),
            item.get("pembeli") or "-",
            item.get("platform_jual") or "-",
        ])

    _apply_table_styling(ws, header_row=1)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream

def export_profit(items: List[Dict[str, Any]], summary: Dict[str, Any]) -> BytesIO:
    """Ekspor laporan laba/rugi dengan sheet ringkasan dan detail."""
    wb = Workbook()
    
    # Sheet 1: Ringkasan Profit
    ws_summary = wb.active
    ws_summary.title = "Ringkasan Profit"

    summary_headers = ["Parameter", "Nilai"]
    ws_summary.append(summary_headers)

    summary_rows = [
        ["Total Item Terjual", summary.get("total_item_terjual", summary.get("total_item", 0))],
        ["Total Kuantitas Terjual", summary.get("total_qty", 0)],
        ["Total Omzet / Penjualan", summary.get("total_jual", 0.0)],
        ["Total Biaya HPP", summary.get("total_hpp_cost", summary.get("total_hpp", 0.0))],
        ["Total Keuntungan (Profit)", summary.get("total_profit", 0.0)],
        ["Rata-rata Margin (%)", f"{summary.get('avg_margin', 0.0)}%"],
    ]
    for row in summary_rows:
        ws_summary.append(row)

    _apply_table_styling(ws_summary, header_row=1)

    # Sheet 2: Detail Profit
    ws_detail = wb.create_sheet(title="Detail Profit")
    detail_headers = [
        "Koleksi ID", "Nama Item", "Total Qty Jual",
        "Total Penjualan", "Total Biaya HPP", "Total Profit", "Margin (%)"
    ]
    ws_detail.append(detail_headers)

    for item in items:
        ws_detail.append([
            item.get("koleksi_id"),
            item.get("nama_item"),
            item.get("total_qty_jual", 0),
            item.get("total_jual", 0.0),
            item.get("total_hpp_cost", 0.0),
            item.get("total_profit", 0.0),
            item.get("margin_pct", 0.0),
        ])

    _apply_table_styling(ws_detail, header_row=1)

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream
